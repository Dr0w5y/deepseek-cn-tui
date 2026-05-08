"""
PPT 文档解析器
=============
解析 .pptx 文件，提取所有幻灯片中的文本内容。

用法：
    from 源码.解析器.PPT解析器 import 解析PPT文档
    文本 = 解析PPT文档("演示.pptx")
    print(文本)
"""

from pathlib import Path
from typing import List


def 解析PPT文档(文件路径: str) -> str:
    """
    解析 PowerPoint (.pptx) 文件，提取所有幻灯片的文本内容。

    参数：
        文件路径: PPT 文档的路径（相对或绝对）

    返回：
        格式化后的文本字符串，每页以 "【第N页】" 开头：
            【第1页】
            标题文本
            正文内容...

            【第2页】
            标题文本
            正文内容...

        解析失败时返回以 "❌" 开头的错误信息
    """

    # ── 文件存在性检查 ──────────────────────────────────
    路径 = Path(文件路径)
    if not 路径.exists():
        return f"❌ 文件不存在：{文件路径}"

    if not 路径.is_file():
        return f"❌ 路径不是文件：{文件路径}"

    if 路径.suffix.lower() != ".pptx":
        return f"❌ 不支持的文件格式（{路径.suffix}），请提供 .pptx 文件"

    # ── 解析文档 ────────────────────────────────────────
    try:
        from pptx import Presentation
    except ImportError:
        return "❌ 缺少 python-pptx 依赖，请执行：pip install python-pptx"

    try:
        演示文稿 = Presentation(str(路径))
    except Exception as e:
        return f"❌ 无法打开 PPT 文档：{e}"

    输出行: List[str] = []

    try:
        幻灯片列表 = 演示文稿.slides
    except Exception as e:
        return f"❌ 无法读取幻灯片：{e}"

    总页数 = len(幻灯片列表)

    for 页码, 幻灯片 in enumerate(幻灯片列表, 1):
        if 输出行:
            输出行.append("")  # 幻灯片之间空行分隔

        输出行.append(f"【第{页码}页】")

        文本数量 = 0

        try:
            for 形状 in 幻灯片.shapes:
                # ── 文本框 ──────────────────────────────
                if 形状.has_text_frame:
                    try:
                        for 段落 in 形状.text_frame.paragraphs:
                            文本 = 段落.text.strip()
                            if 文本:
                                输出行.append(文本)
                                文本数量 += 1
                    except Exception:
                        pass

                # ── 表格 ────────────────────────────────
                if 形状.has_table:
                    try:
                        表格 = 形状.table
                        输出行.append("")  # 表格前空行
                        for 行 in 表格.rows:
                            单元格 = [cell.text.replace("\n", " ").strip() for cell in 行.cells]
                            if any(单元格):
                                输出行.append(" | ".join(单元格))
                                文本数量 += 1
                    except Exception:
                        pass

                # ── 组合 ────────────────────────────────
                if 形状.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
                    try:
                        for 子形状 in 形状.shapes:
                            if 子形状.has_text_frame:
                                for 段落 in 子形状.text_frame.paragraphs:
                                    文本 = 段落.text.strip()
                                    if 文本:
                                        输出行.append(文本)
                                        文本数量 += 1
                    except Exception:
                        pass

        except Exception as e:
            输出行.append(f"⚠️ 幻灯片解析异常：{e}")

        if 文本数量 == 0:
            输出行.append("（无文本内容）")

    if not 输出行:
        return "⚠️ 未提取到任何幻灯片文本内容"

    # ── 添加汇总信息 ────────────────────────────────────
    输出行.append("")
    输出行.append(f"（共 {总页数} 页）")

    return "\n".join(输出行)
