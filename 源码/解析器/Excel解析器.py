"""
Excel 文档解析器
===============
解析 .xlsx 文件，提取所有工作表的行数据。

用法：
    from 源码.解析器.Excel解析器 import 解析Excel文档
    文本 = 解析Excel文档("数据.xlsx")
    print(文本)
"""

from pathlib import Path
from typing import List


def 解析Excel文档(文件路径: str) -> str:
    """
    解析 Excel (.xlsx) 文件，提取所有工作表的行数据。

    参数：
        文件路径: Excel 文档的路径（相对或绝对）

    返回：
        格式化后的文本字符串，每个工作表以 "【工作表：名称】" 开头，
        行数据以 " | " 分隔单元格：
            【工作表：Sheet1】
            表头1 | 表头2 | 表头3
            数据1 | 数据2 | 数据3

            【工作表：Sheet2】
            列A | 列B
            值1 | 值2

        解析失败时返回以 "❌" 开头的错误信息
    """

    # ── 文件存在性检查 ──────────────────────────────────
    路径 = Path(文件路径)
    if not 路径.exists():
        return f"❌ 文件不存在：{文件路径}"

    if not 路径.is_file():
        return f"❌ 路径不是文件：{文件路径}"

    if 路径.suffix.lower() not in (".xlsx", ".xlsm"):
        return f"❌ 不支持的文件格式（{路径.suffix}），请提供 .xlsx 文件"

    # ── 解析文档 ────────────────────────────────────────
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "❌ 缺少 openpyxl 依赖，请执行：pip install openpyxl"

    try:
        工作簿 = load_workbook(str(路径), read_only=True, data_only=True)
    except Exception as e:
        return f"❌ 无法打开 Excel 文档：{e}"

    输出行: List[str] = []

    try:
        工作表名称列表 = 工作簿.sheetnames
    except Exception as e:
        return f"❌ 无法读取工作表列表：{e}"

    for 名称 in 工作表名称列表:
        try:
            工作表 = 工作簿[名称]
        except Exception as e:
            输出行.append(f"⚠️ 无法读取工作表「{名称}」：{e}")
            continue

        if 输出行:
            输出行.append("")  # 工作表之间空行分隔

        输出行.append(f"【工作表：{名称}】")

        行计数 = 0
        try:
            for 行 in 工作表.iter_rows(values_only=True):
                # 将每个单元格转为字符串，None 转为空串
                单元格文本 = [str(cell) if cell is not None else "" for cell in 行]
                # 跳过完全为空的行
                if any(单元格文本):
                    输出行.append(" | ".join(单元格文本))
                    行计数 += 1
        except Exception as e:
            输出行.append(f"⚠️ 读取数据异常：{e}")

        if 行计数 == 0:
            输出行.append("（空工作表）")

    try:
        工作簿.close()
    except Exception:
        pass

    if not 输出行:
        return "⚠️ 未提取到任何工作表数据"

    return "\n".join(输出行)
