"""
Word 文档解析器
==============
解析 .docx 文件，提取段落文本和表格内容。

用法：
    from 源码.解析器.Word解析器 import 解析Word文档
    文本 = 解析Word文档("报告.docx")
    print(文本)
"""

from pathlib import Path
from typing import List


def 解析Word文档(文件路径: str) -> str:
    """
    解析 Word (.docx) 文件，提取所有段落文本和表格内容。

    参数：
        文件路径: Word 文档的路径（相对或绝对）

    返回：
        格式化后的文本字符串，结构为：
            【段落】
            段落1
            段落2

            【表格】
            表头1 | 表头2
            数据1 | 数据2

        解析失败时返回以 "❌" 开头的错误信息
    """

    # ── 文件存在性检查 ──────────────────────────────────
    路径 = Path(文件路径)
    if not 路径.exists():
        return f"❌ 文件不存在：{文件路径}"

    if not 路径.is_file():
        return f"❌ 路径不是文件：{文件路径}"

    if 路径.suffix.lower() != ".docx":
        return f"❌ 不支持的文件格式（{路径.suffix}），请提供 .docx 文件"

    # ── 解析文档 ────────────────────────────────────────
    try:
        from docx import Document
    except ImportError:
        return "❌ 缺少 python-docx 依赖，请执行：pip install python-docx"

    try:
        文档 = Document(str(路径))
    except Exception as e:
        return f"❌ 无法打开 Word 文档：{e}"

    输出行: List[str] = []

    # ── 提取段落 ────────────────────────────────────────
    try:
        段落列表 = [p.text for p in 文档.paragraphs if p.text.strip()]
    except Exception as e:
        段落列表 = []
        输出行.append(f"⚠️ 段落提取异常：{e}")

    if 段落列表:
        输出行.append("【段落】")
        输出行.extend(段落列表)

    # ── 提取表格 ────────────────────────────────────────
    try:
        表格列表 = 文档.tables
    except Exception:
        表格列表 = []

    if 表格列表:
        if 输出行:
            输出行.append("")  # 空行分隔
        输出行.append("【表格】")

        for 表索引, 表格 in enumerate(表格列表, 1):
            if len(表格列表) > 1:
                输出行.append(f"-- 表格 {表索引} --")

            try:
                行数据 = []
                for 行 in 表格.rows:
                    单元格 = [cell.text.replace("\n", " ").strip() for cell in 行.cells]
                    行数据.append(" | ".join(单元格))
                输出行.extend(行数据)
            except Exception as e:
                输出行.append(f"⚠️ 表格 {表索引} 解析异常：{e}")

    if not 输出行:
        return "⚠️ 文档中未提取到任何内容（段落和表格均为空）"

    return "\n".join(输出行)
